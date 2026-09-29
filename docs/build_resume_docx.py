from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT = PROJECT_ROOT.parent / "resume.docx"
BLACK = RGBColor(0x11, 0x11, 0x11)
MUTED = RGBColor(0x55, 0x55, 0x55)


def set_font(run, name: str = "Microsoft YaHei", size: float = 10, bold: bool = False, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def set_cell_shading(cell, fill: str):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def remove_paragraph_borders(style) -> None:
    paragraph_properties = style._element.get_or_add_pPr()
    borders = paragraph_properties.find(qn("w:pBdr"))
    if borders is not None:
        paragraph_properties.remove(borders)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(1)
    normal.paragraph_format.line_spacing = 1.0

    for style_name, size in (("Title", 22), ("Heading 1", 10.5)):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        remove_paragraph_borders(style)

    if "Resume Small" not in [style.name for style in document.styles]:
        style = document.styles.add_style("Resume Small", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
    else:
        style = document.styles["Resume Small"]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(9.1)
    style.font.color.rgb = MUTED
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.0


def add_section_heading(document: Document, text: str):
    paragraph = document.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_font(run, size=10.5, bold=True)
    return paragraph


def add_bullet(document: Document, parts: list[tuple[str, bool]]):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.14)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing = 0.98
    bullet = paragraph.add_run("• ")
    set_font(bullet, size=9.1)
    for text, bold in parts:
        run = paragraph.add_run(text)
        set_font(run, size=9.1, bold=bold)
    return paragraph


def add_body(document: Document, text: str):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    set_font(run, size=9.2)
    return paragraph


def add_project_header(document: Document, title: str, role: str, period: str):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(7.28), WD_TAB_ALIGNMENT.RIGHT)
    left = paragraph.add_run(title)
    set_font(left, size=10.4, bold=True)
    right = paragraph.add_run(f"\t{role} ｜ {period}")
    set_font(right, size=8.8, color=MUTED)
    return paragraph


def build() -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.35)
    section.bottom_margin = Inches(0.22)
    section.left_margin = Inches(0.58)
    section.right_margin = Inches(0.58)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.12)
    configure_styles(document)

    document.core_properties.author = ""
    document.core_properties.title = "AI Agent 工程师简历"
    document.core_properties.subject = "AI Agent 工程师简历"

    title = document.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(0)
    title.paragraph_format.keep_with_next = True
    title_run = title.add_run("[姓名]")
    set_font(title_run, size=23, bold=True)

    role = document.add_paragraph()
    role.paragraph_format.space_after = Pt(1)
    role.paragraph_format.keep_with_next = True
    role_run = role.add_run("AI Agent 工程师 / LLM 应用工程师")
    set_font(role_run, size=11.5, bold=True)

    contact = document.add_paragraph(style="Resume Small")
    contact.paragraph_format.space_after = Pt(4)
    contact.paragraph_format.keep_with_next = True
    contact_run = contact.add_run("手机：[待补充]  ｜  邮箱：[待补充]  ｜  城市：深圳  ｜  GitHub：[待补充]")
    set_font(contact_run, size=8.8, color=MUTED)

    add_section_heading(document, "个人概述")
    add_body(document, "熟悉 Python、Rust 后端开发与 AI 应用工程实践，完成事件驱动算法做市系统及两类 RAG Agent 项目，能够围绕异步并发、状态一致性、风险控制、LangGraph、MCP、检索、评测和容器化完成方案设计与工程验证。")

    add_section_heading(document, "技术能力")
    add_bullet(document, [("编程与后端：", True), ("熟悉 Python、Rust；掌握 FastAPI、Tokio、异步编程、HTTP API、SQLite、PostgreSQL、Docker Compose", False)])
    add_bullet(document, [("Agent 工程：", True), ("LangGraph、MCP Python SDK、Prompt、Function Calling、Tool Use、Memory", False)])
    add_bullet(document, [("AI 工具与开发流程：", True), ("熟练使用 Codex、Claude Code 完成分析、开发、测试与 Review；熟悉 Vibe Coding 工作流", False)])
    add_bullet(document, [("可靠性与工程化：", True), ("模型、检索和工具调用超时控制、结构化错误处理、降级回答、Docker Compose、可观测性", False)])
    add_bullet(document, [("RAG 与检索：", True), ("SQLite FTS5、BM25、FastEmbed、pgvector、向量检索、混合检索、RRF、中文双字词检索、citation", False)])
    add_bullet(document, [("模型与质量：", True), ("OpenAI-compatible API、结构化 JSON 输出、Langfuse、DeepEval、回归评测设计", False)])

    add_section_heading(document, "项目经历")
    add_project_header(document, "Rust 事件驱动算法做市与风险控制系统", "Rust 后端 / 量化系统工程 / 独立开发", "2026.08 - 至今")
    add_body(document, "面向实时行情与模拟交易场景，构建可回放、可审计、可安全退出的事件驱动做市系统。")
    add_bullet(document, [("策略与状态：", True), ("基于 Rust/Tokio 实现 Grid 与 GLFT 两类做市策略，接入交易所级 REST/WebSocket 行情与账户事件，支持纸面、测试网和实盘执行适配；设计有界队列、断线重连、状态回补、订单幂等去重和策略订单归属，保证行情、成交、仓位与订单状态一致。", False)])
    add_bullet(document, [("风险与可靠性：", True), ("实现库存、回撤、盘口过期、价格跳变、滑点、仓位分歧和连续失败等风控，并加入限流退避、死手开关、watchdog 和自动撤单/平仓。", False)])
    add_bullet(document, [("审计与验证：", True), ("使用 SQLite 记录订单版本、成交、markout 和 PnL 快照，接入 Prometheus 指标；补充纸面部分成交与队列模型、参数扫描和样本外评估流程，Rust 测试共 245 个。", False)])

    add_project_header(document, "AI 研发知识库与工单 Agent", "AI Agent 工程 / 独立开发", "2026.09 - 至今")
    add_body(document, "面向研发文档、代码和工单查询场景，构建轻量、可观测、可评测的 Agent 应用。")
    add_bullet(document, [("状态编排：", True), ("基于 LangGraph 拆分路由、知识检索、MCP 工具调用和回答节点，支持 knowledge / repo / issue / direct 意图分流及 thread_id 会话状态。", False)])
    add_bullet(document, [("工具协议：", True), ("基于 MCP Python SDK 实现工具发现和调用，提供知识检索、索引重建、代码搜索和工单查询工具；统一处理模型、检索和工具调用的超时、失败与降级回答。", False)])
    add_bullet(document, [("本地 RAG：", True), ("基于 SQLite FTS5 支持 Markdown、文本及常见代码文件，实现 1200 字符切块、150 字符重叠、BM25 排序、中文双字词检索和来源行号记录。", False)])
    add_bullet(document, [("工程化交付：", True), ("使用 FastAPI 提供 /chat、/health 接口，通过 Docker Compose 编排 API 与 MCP 服务；补充中文检索单元测试和 MCP 网络健康检查。", False)])
    add_bullet(document, [("可观测与评测：", True), ("支持可选 Langfuse 节点追踪，提供 DeepEval 回归评测脚本，为答案相关性、Faithfulness、引用质量和工具选择评估预留统一入口。", False)])

    add_project_header(document, "Enterprise RAG Knowledge Platform", "RAG / AI Agent 工程 / 独立开发", "2026.09 - 至今")
    add_body(document, "面向企业文档、工单和代码知识库场景，搭建可从本地 SQLite 平滑演进到 PostgreSQL 的 RAG 平台。")
    add_bullet(document, [("数据与检索：", True), ("设计文档导入、切分、元数据保存和引用返回链路，使用 FastEmbed 生成向量，结合 PostgreSQL tsvector、pgvector 和 RRF 融合，支持回答溯源与权限控制扩展。", False)])
    add_bullet(document, [("Agent 与部署：", True), ("基于 LangGraph 编排检索与回答流程，使用 MCP 暴露知识检索、文档导入和工单查询工具；通过 FastAPI 与 Docker Compose 完成 PostgreSQL、pgvector、FastEmbed 和 Redis 环境验证。", False)])
    add_bullet(document, [("质量验证：", True), ("补充 4 个单元测试，并完成真实文档导入、混合检索和带引用回答验证。", False)])

    compact = document.add_paragraph(style="Resume Small")
    compact.paragraph_format.space_before = Pt(3)
    compact.paragraph_format.space_after = Pt(0)
    compact.paragraph_format.line_spacing = 0.95
    compact_run = compact.add_run("教育：")
    set_font(compact_run, size=8.4, bold=True)
    compact_run = compact.add_run("[学校] ｜ [专业] ｜ [学历] ｜ [年份-年份]    ")
    set_font(compact_run, size=8.4, color=MUTED)
    compact_run = compact.add_run("求职：")
    set_font(compact_run, size=8.4, bold=True)
    compact_run = compact.add_run("AI Agent / LLM 应用工程师")
    set_font(compact_run, size=8.4, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("AI Agent 工程师简历")
    set_font(footer_run, size=8, color=MUTED)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
